import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PPTX_OUTPUT = r"c:\Users\user\OneDrive - Sri Lanka Institute of Information Technology\Desktop\TNT-SuperMarket\Online-TNT-Supermarket-BE\docs\SE3112_JMeter_Performance_Slides.pptx"
SCREENSHOTS_DIR = r"c:\Users\user\OneDrive - Sri Lanka Institute of Information Technology\Desktop\TNT-SuperMarket\Online-TNT-Supermarket-BE\Performance-Screenshots"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6]

DARK_BG = RGBColor(15, 17, 26)
CARD_BG = RGBColor(26, 30, 46)
TEXT_WHITE = RGBColor(255, 255, 255)
TEXT_MUTED = RGBColor(160, 165, 185)
ACCENT_PURPLE = RGBColor(124, 93, 250)
ACCENT_GREEN = RGBColor(16, 185, 129)
ACCENT_CYAN = RGBColor(56, 189, 248)
BORDER_GRAY = RGBColor(45, 52, 75)
CODE_BG = RGBColor(10, 12, 18)

def apply_background(slide, color):
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = color
    bg_shape.line.fill.background()
    return bg_shape

def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_GRAY):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card

# ==========================================
# SLIDE 1: Approach, Setup & Execution
# ==========================================
slide1 = prs.slides.add_slide(blank_slide_layout)
apply_background(slide1, DARK_BG)

# Header
tb_h = slide1.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.5), Inches(1.1))
tf_h = tb_h.text_frame
tf_h.word_wrap = True
tf_h.margin_left = tf_h.margin_top = tf_h.margin_right = tf_h.margin_bottom = 0

p1 = tf_h.paragraphs[0]
p1.text = "Load & Performance Testing — Apache JMeter"
p1.font.name = "Calibri"
p1.font.size = Pt(28)
p1.font.bold = True
p1.font.color.rgb = TEXT_WHITE

p2 = tf_h.add_paragraph()
p2.text = "SE3112 Advanced Software Engineering  |  Individual Assessment  |  TNT Online Supermarket"
p2.font.name = "Calibri"
p2.font.size = Pt(13)
p2.font.color.rgb = TEXT_MUTED

# Student Badge Right
badge1 = create_card(slide1, Inches(10.6), Inches(0.45), Inches(1.9), Inches(0.7), bg_color=ACCENT_PURPLE, border_color=None)
tf_b1 = badge1.text_frame
tf_b1.word_wrap = True
p_b1 = tf_b1.paragraphs[0]
p_b1.alignment = PP_ALIGN.CENTER
p_b1.text = "ID: IT24101022\nSlide 1 of 2"
p_b1.font.name = "Calibri"
p_b1.font.size = Pt(11)
p_b1.font.bold = True
p_b1.font.color.rgb = TEXT_WHITE

# Card 1: Objective & Workload Model (Left Column)
card_left1 = create_card(slide1, Inches(0.8), Inches(1.6), Inches(5.8), Inches(2.9))
tf_c1 = card_left1.text_frame
tf_c1.word_wrap = True
tf_c1.margin_left = Inches(0.25)
tf_c1.margin_right = Inches(0.25)
tf_c1.margin_top = Inches(0.2)

p = tf_c1.paragraphs[0]
p.text = "TARGET ARCHITECTURE & WORKLOAD MODEL"
p.font.name = "Calibri"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_PURPLE

bullets1 = [
    "Objective: Simulate 110 concurrent active shoppers hitting Ocelot API Gateway.",
    "Scope: Test resilience across Identity, Catalog, Order & Reporting microservices.",
    "Workload Parameters: 110 Virtual Users, 10s ramp-up period, 5 loops per user.",
    "Human Pacing: 100ms - 300ms Uniform Random Timer (avoids socket starvation).",
    "Dual-Layer Assertions: HTTP 200 OK Status and strict SLA Latency < 1,500ms."
]
for b in bullets1:
    bp = tf_c1.add_paragraph()
    bp.text = "• " + b
    bp.font.name = "Calibri"
    bp.font.size = Pt(11.5)
    bp.font.color.rgb = TEXT_MUTED

# Card 2: 8 Endpoints Matrix (Below Card 1)
card_ep = create_card(slide1, Inches(0.8), Inches(4.65), Inches(5.8), Inches(1.15))
tf_ep = card_ep.text_frame
tf_ep.word_wrap = True
tf_ep.margin_left = Inches(0.25)
tf_ep.margin_right = Inches(0.25)
tf_ep.margin_top = Inches(0.15)

p = tf_ep.paragraphs[0]
p.text = "8 MISSION-CRITICAL ENDPOINTS TESTED"
p.font.name = "Calibri"
p.font.size = Pt(10)
p.font.bold = True
p.font.color.rgb = ACCENT_CYAN

p_ep_list = tf_ep.add_paragraph()
p_ep_list.text = "Gateway /health  •  Identity /health  •  Catalog /health  •  Browse /products\nBrowse /categories  •  Order /health  •  Reporting /health  •  Token Filter /products"
p_ep_list.font.name = "Calibri"
p_ep_list.font.size = Pt(10.5)
p_ep_list.font.color.rgb = TEXT_WHITE

# Card 3: Execution Commands Box (Below Endpoints)
card_cmd = create_card(slide1, Inches(0.8), Inches(5.95), Inches(5.8), Inches(1.1), bg_color=CODE_BG)
tf_cmd = card_cmd.text_frame
tf_cmd.word_wrap = True
tf_cmd.margin_left = Inches(0.25)
tf_cmd.margin_top = Inches(0.12)

p_cmd_title = tf_cmd.paragraphs[0]
p_cmd_title.text = ">_ EXECUTION COMMANDS (HEADLESS NON-GUI)"
p_cmd_title.font.name = "Consolas"
p_cmd_title.font.size = Pt(9.5)
p_cmd_title.font.bold = True
p_cmd_title.font.color.rgb = ACCENT_GREEN

p_cmd_body = tf_cmd.add_paragraph()
p_cmd_body.text = "jmeter -n -t TNT_Supermarket_Performance_Test.jmx -l jmeter_results.jtl -e -o jmeter_html_report\nOne-Click Batch Runner: .\\run_performance_test.bat"
p_cmd_body.font.name = "Consolas"
p_cmd_body.font.size = Pt(9.5)
p_cmd_body.font.color.rgb = RGBColor(220, 225, 240)

# Right Column: Screenshot Box
img1_path = os.path.join(SCREENSHOTS_DIR, "05_jmeter_response_time_percentiles.png")
card_img1 = create_card(slide1, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.45))

if os.path.exists(img1_path):
    slide1.shapes.add_picture(img1_path, Inches(7.05), Inches(1.8), Inches(5.3), Inches(4.7))

tb_cap1 = slide1.shapes.add_textbox(Inches(6.9), Inches(6.6), Inches(5.6), Inches(0.4))
tf_cap1 = tb_cap1.text_frame
tf_cap1.margin_left = tf_cap1.margin_top = tf_cap1.margin_right = tf_cap1.margin_bottom = 0
p_cap1 = tf_cap1.paragraphs[0]
p_cap1.alignment = PP_ALIGN.CENTER
p_cap1.text = "Figure: Apache JMeter Response Time Percentiles (95% requests < 28ms, SLA < 1500ms)"
p_cap1.font.name = "Calibri"
p_cap1.font.size = Pt(10)
p_cap1.font.italic = True
p_cap1.font.color.rgb = TEXT_MUTED


# ==========================================
# SLIDE 2: Results & GitHub Actions CI/CD Proof
# ==========================================
slide2 = prs.slides.add_slide(blank_slide_layout)
apply_background(slide2, DARK_BG)

# Header
tb_h2 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.5), Inches(1.1))
tf_h2 = tb_h2.text_frame
tf_h2.word_wrap = True
tf_h2.margin_left = tf_h2.margin_top = tf_h2.margin_right = tf_h2.margin_bottom = 0

p1 = tf_h2.paragraphs[0]
p1.text = "Benchmark Results & GitHub Actions CI/CD Proof"
p1.font.name = "Calibri"
p1.font.size = Pt(28)
p1.font.bold = True
p1.font.color.rgb = TEXT_WHITE

p2 = tf_h2.add_paragraph()
p2.text = "110 Virtual Users Peak Load  |  Branch: feature/se3112-performance-jmeter  |  Author: IT24101022"
p2.font.name = "Calibri"
p2.font.size = Pt(13)
p2.font.color.rgb = TEXT_MUTED

# Student Badge Right
badge2 = create_card(slide2, Inches(10.6), Inches(0.45), Inches(1.9), Inches(0.7), bg_color=ACCENT_GREEN, border_color=None)
tf_b2 = badge2.text_frame
tf_b2.word_wrap = True
p_b2 = tf_b2.paragraphs[0]
p_b2.alignment = PP_ALIGN.CENTER
p_b2.text = "ID: IT24101022\nSlide 2 of 2"
p_b2.font.name = "Calibri"
p_b2.font.size = Pt(11)
p_b2.font.bold = True
p_b2.font.color.rgb = TEXT_WHITE

# Left Column: Metrics Cards
m_w = Inches(1.8)
m_h = Inches(1.1)

# Metric 1: Throughput
m1 = create_card(slide2, Inches(0.8), Inches(1.6), m_w, m_h, bg_color=CARD_BG)
tf_m1 = m1.text_frame
tf_m1.word_wrap = True
p = tf_m1.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
p.text = "440.97"
p.font.name = "Calibri"
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p_sub = tf_m1.add_paragraph()
p_sub.alignment = PP_ALIGN.CENTER
p_sub.text = "Throughput (req/s)"
p_sub.font.size = Pt(9.5)
p_sub.font.color.rgb = TEXT_MUTED

# Metric 2: Median Latency
m2 = create_card(slide2, Inches(2.75), Inches(1.6), m_w, m_h, bg_color=CARD_BG)
tf_m2 = m2.text_frame
tf_m2.word_wrap = True
p = tf_m2.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
p.text = "7.00 ms"
p.font.name = "Calibri"
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = ACCENT_CYAN
p_sub = tf_m2.add_paragraph()
p_sub.alignment = PP_ALIGN.CENTER
p_sub.text = "Median Latency (p50)"
p_sub.font.size = Pt(9.5)
p_sub.font.color.rgb = TEXT_MUTED

# Metric 3: Error Rate
m3 = create_card(slide2, Inches(4.7), Inches(1.6), m_w, m_h, bg_color=CARD_BG)
tf_m3 = m3.text_frame
tf_m3.word_wrap = True
p = tf_m3.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
p.text = "0.00%"
p.font.name = "Calibri"
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p_sub = tf_m3.add_paragraph()
p_sub.alignment = PP_ALIGN.CENTER
p_sub.text = "Error Rate (4,400 Pass)"
p_sub.font.size = Pt(9.5)
p_sub.font.color.rgb = TEXT_MUTED

# SLA Compliance Card (Left Column)
card_sla = create_card(slide2, Inches(0.8), Inches(2.85), Inches(5.7), Inches(2.55))
tf_sla = card_sla.text_frame
tf_sla.word_wrap = True
tf_sla.margin_left = Inches(0.25)
tf_sla.margin_top = Inches(0.18)

p = tf_sla.paragraphs[0]
p.text = "SLA COMPLIANCE & PERFORMANCE METRICS"
p.font.name = "Calibri"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

sla_items = [
    "Total Samples: 4,400 Requests Dispatched (100% Passed without timeouts)",
    "SLA Threshold: Strict < 1,500ms Assertion (Zero SLA Breaches)",
    "95th Percentile (p95): 28.00 ms (53× Faster than Required SLA)",
    "99th Percentile (p99): 60.00 ms under sustained 110 user concurrent load",
    "APDEX Score: 0.98+ / 1.00 (Ranked 'Excellent' User Satisfaction)"
]
for item in sla_items:
    p_i = tf_sla.add_paragraph()
    p_i.text = "✓ " + item
    p_i.font.name = "Calibri"
    p_i.font.size = Pt(11)
    p_i.font.color.rgb = TEXT_WHITE

# Pipeline Summary Card (Below SLA)
card_pipe_sum = create_card(slide2, Inches(0.8), Inches(5.55), Inches(5.7), Inches(1.5), bg_color=CARD_BG)
tf_psum = card_pipe_sum.text_frame
tf_psum.word_wrap = True
tf_psum.margin_left = Inches(0.25)
tf_psum.margin_top = Inches(0.15)

p = tf_psum.paragraphs[0]
p.text = "CI/CD PIPELINE AUTOMATION HIGHLIGHTS"
p.font.name = "Calibri"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = ACCENT_CYAN

p_ptext = tf_psum.add_paragraph()
p_ptext.text = "• Standalone 3-Stage Pipeline: Validate -> JMeter Load Test -> Summary Table.\n• Fast Feedback: Total workflow runtime is only 2m 1s with 100% green checks.\n• Downloadable Artifacts: Complete interactive HTML dashboard automatically archived."
p_ptext.font.name = "Calibri"
p_ptext.font.size = Pt(10.5)
p_ptext.font.color.rgb = TEXT_MUTED

# Right Column: Two Screenshots (GitHub Actions Run + Commits)
img_act_path = os.path.join(SCREENSHOTS_DIR, "06_github_actions_jmeter_pipeline.png")
img_com_path = os.path.join(SCREENSHOTS_DIR, "07_github_commits_it24101022.png")

# Box 1: Actions Pipeline Screenshot
card_act = create_card(slide2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(3.2))
if os.path.exists(img_act_path):
    slide2.shapes.add_picture(img_act_path, Inches(6.9), Inches(1.7), Inches(5.5), Inches(2.75))

tb_cap_act = slide2.shapes.add_textbox(Inches(6.8), Inches(4.5), Inches(5.7), Inches(0.3))
tf_cap_act = tb_cap_act.text_frame
tf_cap_act.margin_left = tf_cap_act.margin_top = tf_cap_act.margin_bottom = tf_cap_act.margin_right = 0
p = tf_cap_act.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
p.text = "Figure 1: GitHub Actions Standalone Pipeline — 3-Stage Run by IT24101022 (Success: 2m 1s)"
p.font.name = "Calibri"
p.font.size = Pt(9.5)
p.font.italic = True
p.font.color.rgb = TEXT_MUTED

# Box 2: Commits History Screenshot
card_com = create_card(slide2, Inches(6.8), Inches(4.9), Inches(5.7), Inches(2.15))
if os.path.exists(img_com_path):
    slide2.shapes.add_picture(img_com_path, Inches(6.9), Inches(5.0), Inches(5.5), Inches(1.75))

tb_cap_com = slide2.shapes.add_textbox(Inches(6.8), Inches(6.8), Inches(5.7), Inches(0.25))
tf_cap_com = tb_cap_com.text_frame
tf_cap_com.margin_left = tf_cap_com.margin_top = tf_cap_com.margin_bottom = tf_cap_com.margin_right = 0
p = tf_cap_com.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
p.text = "Figure 2: Verified Commit History on feature/se3112-performance-jmeter by IT24101022"
p.font.name = "Calibri"
p.font.size = Pt(9.5)
p.font.italic = True
p.font.color.rgb = TEXT_MUTED

prs.save(PPTX_OUTPUT)
print(f"PowerPoint Presentation successfully created at: {PPTX_OUTPUT}")
